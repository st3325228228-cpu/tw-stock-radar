"""
台股技術分析儀表板｜精美 UI 版
執行：streamlit run app.py
"""
import html as htmllib
import re
import time
import warnings
from datetime import datetime, timedelta
from io import BytesIO, StringIO
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import requests
import streamlit as st
import yfinance as yf
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from plotly.subplots import make_subplots

warnings.filterwarnings("ignore", category=FutureWarning)

st.set_page_config(page_title="TW Stock Radar｜台股技術分析", page_icon="📈",
                   layout="wide", initial_sidebar_state="expanded")

# ══════════════════════════════════════════════════════════════════════════════
# 0. 設定 & 色彩
# ══════════════════════════════════════════════════════════════════════════════
TZ = ZoneInfo("Asia/Taipei")
DEFAULT = {
    "SHEET_ID":     "1Z8_YFX2MrkDJfbfRMNlsbIZK0lUPj-GCdiVN0DCtx1E",
    "GID":          "1522695003",
    "TICKER_COL":   "TICKERS",
    "DISPLAY_DAYS": 91,
    "WARMUP_DAYS":  400,
    "RETRY":        3,
}
HEADERS = {"User-Agent": "Mozilla/5.0"}

UP, DOWN = "#FF4D4F", "#22C55E"          # 台股：紅漲綠跌
ACCENT, ACCENT2 = "#6C8CFF", "#A855F7"
MUTED = "#8B93AD"
SCORE_RANGE = 18                          # 理論最大分數約 ±18

# ══════════════════════════════════════════════════════════════════════════════
# 1. 全域 CSS
# ══════════════════════════════════════════════════════════════════════════════
CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Noto+Sans+TC:wght@400;500;700;900&display=swap');
.stApp{font-family:'Noto Sans TC',sans-serif;
  background:
    radial-gradient(1200px 600px at 5% -10%, rgba(108,140,255,.18), transparent 60%),
    radial-gradient(900px 500px at 100% 0%, rgba(236,72,153,.10), transparent 60%),
    #0B1020;}
header[data-testid="stHeader"]{background:transparent;}
#MainMenu, footer{visibility:hidden;}
.block-container{padding-top:2.2rem;max-width:1440px;}
section[data-testid="stSidebar"]{background:linear-gradient(180deg,#10182D 0%,#0B1020 100%);
  border-right:1px solid rgba(255,255,255,.06);}

/* Brand */
.brand{display:flex;align-items:center;gap:12px;padding:6px 2px 18px;border-bottom:1px solid rgba(255,255,255,.06);margin-bottom:14px}
.brand .logo{width:42px;height:42px;border-radius:12px;display:grid;place-items:center;font-size:22px;
  background:linear-gradient(135deg,#6C8CFF,#A855F7);box-shadow:0 6px 18px rgba(108,140,255,.4)}
.brand b{font-size:1.05rem;color:#fff;display:block}
.brand span{font-size:.75rem;color:#8B93AD}

/* Hero */
.hero{position:relative;overflow:hidden;padding:30px 34px;border-radius:22px;margin-bottom:24px;
  background:linear-gradient(135deg,rgba(108,140,255,.28),rgba(168,85,247,.18) 55%,rgba(236,72,153,.14));
  border:1px solid rgba(255,255,255,.09);box-shadow:0 14px 50px rgba(0,0,0,.35)}
.hero:after{content:"";position:absolute;right:-60px;top:-60px;width:260px;height:260px;border-radius:50%;
  background:radial-gradient(circle,rgba(255,255,255,.12),transparent 70%)}
.hero-badge{display:inline-block;font-size:.7rem;letter-spacing:.2em;color:#C7D2FE;
  background:rgba(255,255,255,.07);padding:5px 12px;border-radius:999px;border:1px solid rgba(255,255,255,.12)}
.hero h1{margin:10px 0 6px;font-weight:900;font-size:2.2rem;line-height:1.2;
  background:linear-gradient(90deg,#FFFFFF,#C7D2FE);-webkit-background-clip:text;-webkit-text-fill-color:transparent}
.hero p{color:#B4BCD6;margin:0;font-size:.95rem}

/* Cards */
.card{background:rgba(20,27,45,.72);border:1px solid rgba(255,255,255,.07);border-radius:16px;
  padding:16px 18px;box-shadow:0 6px 24px rgba(0,0,0,.25);backdrop-filter:blur(8px)}
.section-title{font-size:1.25rem;font-weight:800;margin:30px 0 14px;display:flex;align-items:center;gap:10px;color:#F3F4F6}
.section-title:before{content:"";width:6px;height:22px;border-radius:3px;background:linear-gradient(180deg,#6C8CFF,#A855F7)}

/* KPI */
.kpi-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin:4px 0 16px}
.kpi .lbl{font-size:.78rem;color:#8B93AD;letter-spacing:.03em}
.kpi .val{font-size:1.6rem;font-weight:800;color:#F3F4F6;margin-top:4px;line-height:1.2}
.kpi .sub{font-size:.75rem;color:#8B93AD;margin-top:2px}

/* Stock cards */
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(250px,1fr));gap:14px}
.stock-card{transition:transform .18s ease,border-color .18s ease,box-shadow .18s ease}
.stock-card:hover{transform:translateY(-4px);border-color:rgba(108,140,255,.55);box-shadow:0 12px 30px rgba(108,140,255,.18)}
.sc-top{display:flex;justify-content:space-between;align-items:flex-start;gap:8px}
.sc-code{font-weight:800;font-size:1.05rem;color:#E0E7FF}
.sc-name{font-size:.82rem;color:#8B93AD;max-width:130px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.badge{font-size:.72rem;font-weight:700;padding:4px 10px;border-radius:999px;white-space:nowrap}
.sc-price{font-size:1.65rem;font-weight:800;margin-top:10px;color:#fff}
.sc-chg{font-weight:700;font-size:.9rem}
.spark{width:100%;height:50px;margin:8px 0 6px;display:block}
.sc-foot{display:flex;justify-content:space-between;font-size:.78rem;color:#8B93AD;
  border-top:1px dashed rgba(255,255,255,.08);padding-top:8px}
.sc-foot b{color:#E0E7FF}

/* Detail header */
.detail-head{display:flex;justify-content:space-between;align-items:center;flex-wrap:wrap;gap:18px;margin-bottom:14px}
.dh-code{font-size:.85rem;color:#8B93AD;letter-spacing:.08em}
.dh-name{font-size:1.7rem;font-weight:900;color:#fff}
.dh-price{font-size:2.2rem;font-weight:900;color:#fff;line-height:1}
.dh-chg{font-size:1rem;font-weight:700;margin-left:10px}
.dh-right{text-align:right}
.dh-conf{font-size:.8rem;color:#8B93AD;margin-top:6px}

/* Signals */
.sig-title{font-weight:800;margin-bottom:10px;font-size:1rem}
.sig-title.buy{color:#4ADE80}.sig-title.sell{color:#F87171}
.sig{list-style:none;padding:0;margin:0}
.sig li{padding:9px 12px;margin-bottom:7px;border-radius:10px;font-size:.88rem;display:flex;justify-content:space-between;gap:10px;color:#E5E7EB}
.sig.buy li{background:rgba(34,197,94,.10);border-left:3px solid #22C55E}
.sig.sell li{background:rgba(255,77,79,.10);border-left:3px solid #FF4D4F}
.sig li.none{background:rgba(255,255,255,.04);border-left:3px solid #4B5563;color:#8B93AD}
.sig .w{font-weight:800;font-variant-numeric:tabular-nums}

/* Feature (empty state) */
.feat-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:14px}
.feat .ic{font-size:1.6rem}.feat h4{margin:8px 0 4px;color:#fff}.feat p{color:#8B93AD;font-size:.86rem;margin:0}

/* Widgets */
.stTabs [data-baseweb="tab-list"]{gap:6px;border-bottom:1px solid rgba(255,255,255,.08)}
.stTabs [data-baseweb="tab"]{background:rgba(255,255,255,.035);border-radius:10px 10px 0 0;padding:8px 16px}
.stTabs [aria-selected="true"]{background:rgba(108,140,255,.18)!important}
.stButton>button[kind="primary"],.stDownloadButton>button{
  background:linear-gradient(90deg,#6C8CFF,#A855F7)!important;border:0!important;color:#fff!important;
  font-weight:700;border-radius:12px;box-shadow:0 6px 18px rgba(108,140,255,.35)}
.stButton>button{border-radius:12px}
div[data-testid="stExpander"]{border-radius:14px;border:1px solid rgba(255,255,255,.07)}
.footer{color:#6B7280;font-size:.78rem;text-align:center;margin:40px 0 10px}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# 2. 工具函式
# ══════════════════════════════════════════════════════════════════════════════
def esc(s):
    return htmllib.escape(str(s))


def http_get(url, retry=DEFAULT["RETRY"], encoding=None):
    for i in range(retry):
        try:
            r = requests.get(url, headers=HEADERS, timeout=20)
            r.raise_for_status()
            return r.content.decode(encoding, errors="replace") if encoding else r.text
        except Exception:
            time.sleep(2 * (i + 1))
    return None


def to_cell(v):
    if v is None:
        return None
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, (np.floating, float)):
        return None if np.isnan(v) else float(v)
    if isinstance(v, pd.Timestamp):
        return v.date()
    return v


def fmt(v, spec=".1f"):
    return "—" if v is None or (isinstance(v, float) and np.isnan(v)) else format(v, spec)


def safe_title(title, used):
    title = re.sub(r"[\[\]\:\*\?\/\\]", "_", title)[:31]
    base, n = title, 1
    while title in used:
        suffix = f"_{n}"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title

# ══════════════════════════════════════════════════════════════════════════════
# 3. 股票清單 & 名稱
# ══════════════════════════════════════════════════════════════════════════════
@st.cache_data(ttl=24 * 3600, show_spinner=False)
def build_tw_name_map():
    sources = {
        ".TW":  "https://isin.twse.com.tw/isin/C_public.jsp?strMode=2",
        ".TWO": "https://isin.twse.com.tw/isin/C_public.jsp?strMode=4",
    }
    pat = re.compile(r"^([0-9A-Z]{4,6})[\s\u3000]+(.+)$")
    name_map = {}
    for suffix, url in sources.items():
        page = http_get(url, encoding="cp950")
        if not page:
            continue
        try:
            df = pd.read_html(StringIO(page))[0]
        except Exception:
            continue
        for val in df.iloc[:, 0].dropna().astype(str):
            m = pat.match(val.strip())
            if m and m.group(1) not in name_map:
                name_map[m.group(1)] = {"name": m.group(2).strip(), "suffix": suffix}
    return name_map


@st.cache_data(ttl=600, show_spinner=False)
def load_tickers(sheet_id, gid, ticker_col):
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}"
    df = pd.read_csv(url, dtype=str)
    col = next((c for c in df.columns if c.strip().upper() == ticker_col.strip().upper()), df.columns[0])
    vals = df[col].dropna().astype(str).str.strip()
    return list(dict.fromkeys(v for v in vals if v))


def parse_manual(text):
    vals = [v.strip() for v in re.split(r"[\s,，、;；]+", text or "") if v.strip()]
    return list(dict.fromkeys(vals))


def normalize_ticker(raw, name_map):
    t = str(raw).strip().upper()
    if not t or t == "NAN":
        return None
    if t.endswith((".TW", ".TWO")):
        return t
    if re.fullmatch(r"[0-9A-Z]{4,6}", t) and any(ch.isdigit() for ch in t):
        info = name_map.get(t)
        return t + (info["suffix"] if info else ".TW")
    return t


@st.cache_data(ttl=24 * 3600, show_spinner=False)
def fetch_yf_name(ticker):
    try:
        info = yf.Ticker(ticker).info
        return info.get("longName") or info.get("shortName") or ""
    except Exception:
        return ""


def get_company_name(ticker, name_map):
    code = ticker.split(".")[0]
    return name_map[code]["name"] if code in name_map else fetch_yf_name(ticker)

# ══════════════════════════════════════════════════════════════════════════════
# 4. 下載股價
# ══════════════════════════════════════════════════════════════════════════════
def clean_ohlcv(df):
    need = ["Open", "High", "Low", "Close", "Volume"]
    if df is None or df.empty or not set(need).issubset(df.columns):
        return pd.DataFrame()
    df = df[need].dropna(subset=["Close"])
    df = df[df["Volume"].fillna(0) > 0]
    if df.empty:
        return df
    df = df.reset_index()
    df.columns = ["日期", "開盤價", "最高價", "最低價", "收盤價", "成交量"]
    df["日期"] = pd.to_datetime(df["日期"]).dt.tz_localize(None).dt.date
    return df.sort_values("日期").reset_index(drop=True)


def download_all(tickers, start, end):
    data = {}
    for _ in range(DEFAULT["RETRY"]):
        pending = [t for t in tickers if t not in data]
        if not pending:
            break
        try:
            raw = yf.download(pending, start=start, end=end, auto_adjust=True,
                              group_by="ticker", threads=True, progress=False)
        except Exception:
            time.sleep(3)
            continue
        for t in pending:
            try:
                if isinstance(raw.columns, pd.MultiIndex):
                    if t not in raw.columns.get_level_values(0):
                        continue
                    df = raw[t].copy()
                else:
                    df = raw.copy()
                df = clean_ohlcv(df)
                if not df.empty:
                    data[t] = df
            except Exception:
                continue
        time.sleep(1)
    return data


@st.cache_data(ttl=1800, show_spinner=False)
def download_with_fallback(tickers: tuple, start: str, end: str):
    data = download_all(list(tickers), start, end)
    alt_map = {}
    for t in tickers:
        if t not in data:
            if t.endswith(".TWO"):
                alt_map[t[:-4] + ".TW"] = t
            elif t.endswith(".TW"):
                alt_map[t[:-3] + ".TWO"] = t
    if alt_map:
        alt_data = download_all(list(alt_map), start, end)
        for alt, orig in alt_map.items():
            if alt in alt_data:
                data[orig] = alt_data[alt]
    return data

# ══════════════════════════════════════════════════════════════════════════════
# 5. 技術指標（與原版相同）
# ══════════════════════════════════════════════════════════════════════════════
def _rsi(c, p=14):
    d = c.diff()
    ag = d.clip(lower=0).ewm(alpha=1/p, min_periods=p, adjust=False).mean()
    al = (-d.clip(upper=0)).ewm(alpha=1/p, min_periods=p, adjust=False).mean()
    rsi = 100 - 100 / (1 + ag / al.replace(0, np.nan))
    return rsi.where(al != 0, 100.0).where(ag.notna())


def _macd(c, f=12, s=26, sig=9):
    m = c.ewm(span=f, adjust=False).mean() - c.ewm(span=s, adjust=False).mean()
    sl = m.ewm(span=sig, adjust=False).mean()
    return m, sl, m - sl


def _kd(h, l, c, n=9):
    lo, hi = l.rolling(n).min(), h.rolling(n).max()
    rng = hi - lo
    rsv = ((c - lo) / rng.replace(0, np.nan) * 100).where(rng != 0, 50.0).where(lo.notna())
    K = np.full(len(c), np.nan)
    D = np.full(len(c), np.nan)
    k_prev = d_prev = 50.0
    for i, r in enumerate(rsv.values):
        if np.isnan(r):
            continue
        k_prev = k_prev * 2/3 + r / 3
        d_prev = d_prev * 2/3 + k_prev / 3
        K[i], D[i] = k_prev, d_prev
    return pd.Series(K, index=c.index), pd.Series(D, index=c.index)


def _true_range(h, l, c):
    return pd.concat([h - l, (h - c.shift()).abs(), (l - c.shift()).abs()], axis=1).max(axis=1)


def _atr(h, l, c, p=14):
    return _true_range(h, l, c).ewm(alpha=1/p, min_periods=p, adjust=False).mean()


def _adx(h, l, c, p=14):
    tr = _true_range(h, l, c)
    up, dn = h.diff(), -l.diff()
    pdm = pd.Series(np.where((up > dn) & (up > 0), up, 0.0), index=c.index)
    ndm = pd.Series(np.where((dn > up) & (dn > 0), dn, 0.0), index=c.index)
    sm = lambda x: x.ewm(alpha=1/p, min_periods=p, adjust=False).mean()
    atr_s = sm(tr).replace(0, np.nan)
    pdi, ndi = 100 * sm(pdm) / atr_s, 100 * sm(ndm) / atr_s
    dx = (abs(pdi - ndi) / (pdi + ndi).replace(0, np.nan) * 100)
    return sm(dx), pdi, ndi


def _cci(h, l, c, p=20):
    tp = (h + l + c) / 3
    mad = tp.rolling(p).apply(lambda x: np.abs(x - x.mean()).mean(), raw=True)
    return (tp - tp.rolling(p).mean()) / (0.015 * mad.replace(0, np.nan))


def _sar(h, l, step=0.02, mx=0.2):
    hv, lv = h.values.astype(float), l.values.astype(float)
    n = len(hv)
    sar = np.zeros(n)
    if n == 0:
        return pd.Series(sar, index=h.index)
    bull, sar[0], ep, af = True, lv[0], hv[0], step
    for i in range(1, n):
        prev = sar[i-1] + af * (ep - sar[i-1])
        if bull:
            prev = min(prev, lv[i-1], lv[max(0, i-2)])
            if lv[i] < prev:
                bull, sar[i], ep, af = False, ep, lv[i], step
            else:
                sar[i] = prev
                if hv[i] > ep:
                    ep, af = hv[i], min(af + step, mx)
        else:
            prev = max(prev, hv[i-1], hv[max(0, i-2)])
            if hv[i] > prev:
                bull, sar[i], ep, af = True, ep, hv[i], step
            else:
                sar[i] = prev
                if lv[i] < ep:
                    ep, af = lv[i], min(af + step, mx)
    return pd.Series(sar, index=h.index)


def compute_indicators(df):
    c, h, l, v = (df[k].astype(float) for k in ["收盤價", "最高價", "最低價", "成交量"])
    ind = pd.DataFrame({"日期": df["日期"], "收盤價": c})
    ind["RSI(14)"] = _rsi(c).round(2)
    ml, ms, mh = _macd(c)
    ind["MACD"], ind["MACD_Signal"], ind["MACD_Hist"] = ml.round(4), ms.round(4), mh.round(4)
    for p in (5, 20, 60):
        ind[f"SMA{p}"] = c.rolling(p).mean().round(2)
    ind["SMA50"] = c.rolling(50).mean().round(2)
    ind["SMA200"] = c.rolling(200).mean().round(2)
    ind["EMA12"] = c.ewm(span=12, adjust=False).mean().round(2)
    ind["EMA26"] = c.ewm(span=26, adjust=False).mean().round(2)
    K, D = _kd(h, l, c)
    ind["K(9)"], ind["D(3)"] = K.round(2), D.round(2)
    ind["OBV"] = (np.sign(c.diff()).fillna(0) * v).cumsum().round(0)
    ind["量比(20)"] = (v / v.rolling(20).mean()).round(2)
    bm, sd = c.rolling(20).mean(), c.rolling(20).std()
    bu, bl = bm + 2 * sd, bm - 2 * sd
    ind["BB_Upper"], ind["BB_Mid"], ind["BB_Lower"] = bu.round(2), bm.round(2), bl.round(2)
    ind["BBW"] = ((bu - bl) / bm).round(4)
    ind["%B"] = ((c - bl) / (bu - bl).replace(0, np.nan)).round(4)
    ind["BIAS(20)"] = ((c - bm) / bm * 100).round(2)
    ind["ROC(12)"] = (c.pct_change(12) * 100).round(2)
    ind["ATR(14)"] = _atr(h, l, c).round(4)
    adx, pdi, ndi = _adx(h, l, c)
    ind["ADX(14)"], ind["+DI"], ind["-DI"] = adx.round(2), pdi.round(2), ndi.round(2)
    ind["CCI(20)"] = _cci(h, l, c).round(2)
    ind["Momentum(10)"] = (c - c.shift(10)).round(2)
    tp = (h + l + c) / 3
    ind["VWAP(20)"] = ((tp * v).rolling(20).sum() / v.rolling(20).sum()).round(2)
    ind["SAR"] = _sar(h, l).round(2)
    tk = (h.rolling(9).max() + l.rolling(9).min()) / 2
    kj = (h.rolling(26).max() + l.rolling(26).min()) / 2
    ind["Ichi_Tenkan"], ind["Ichi_Kijun"] = tk.round(2), kj.round(2)
    ind["Ichi_SenkouA"] = ((tk + kj) / 2).shift(26).round(2)
    ind["Ichi_SenkouB"] = ((h.rolling(52).max() + l.rolling(52).min()) / 2).shift(26).round(2)
    ind["Envelope_U"], ind["Envelope_L"] = (bm * 1.025).round(2), (bm * 0.975).round(2)
    s50, s200 = ind["SMA50"], ind["SMA200"]
    ind["Golden_Cross"] = np.select(
        [s50.notna() & s200.notna() & (s50 > s200), s50.notna() & s200.notna() & (s50 < s200)],
        ["黃金交叉(多)", "死亡交叉(空)"], default="N/A")
    return ind

# ══════════════════════════════════════════════════════════════════════════════
# 6. 投資建議（與原版相同）
# ══════════════════════════════════════════════════════════════════════════════
LEVELS = [  # (最低分數, 建議, 底色, 字色, 粗體, emoji)
    (7,    "強力買進 ★★★", "00B050", "FFFFFF", True,  "🚀"),
    (4,    "買進 ★★",       "92D050", "000000", True,  "🟢"),
    (1,    "觀望偏多 ★",    "E2EFDA", "000000", False, "🌱"),
    (0,    "觀望（中性）",   "F2F2F2", "000000", False, "⚪"),
    (-3,   "觀望偏空 ▼",    "FFDDC1", "000000", False, "🟠"),
    (-6,   "賣出 ▼▼",       "FF7070", "FFFFFF", True,  "🔴"),
    (-999, "強力賣出 ▼▼▼", "C00000", "FFFFFF", True,  "⛔"),
]


def score_to_level(score):
    for lv in LEVELS:
        if score >= lv[0]:
            return lv
    return LEVELS[-1]


def generate_advice(df_ind, df_disp):
    valid = df_ind.dropna(subset=["RSI(14)", "MACD_Hist", "K(9)"])
    if len(valid) < 2:
        return None
    cur, prv = valid.iloc[-1], valid.iloc[-2]
    g = lambda row, col: float(row[col]) if pd.notna(row.get(col)) else np.nan
    ok = lambda *xs: all(not np.isnan(x) for x in xs)

    buys, sells = [], []
    def add(w, txt):
        (buys if w > 0 else sells).append((w, txt))

    close, p_close = g(cur, "收盤價"), g(prv, "收盤價")
    rsi, p_rsi = g(cur, "RSI(14)"), g(prv, "RSI(14)")
    h, p_h = g(cur, "MACD_Hist"), g(prv, "MACD_Hist")
    k, d, pk, pd_ = g(cur, "K(9)"), g(cur, "D(3)"), g(prv, "K(9)"), g(prv, "D(3)")
    pct_b, cci = g(cur, "%B"), g(cur, "CCI(20)")
    sar, p_sar = g(cur, "SAR"), g(prv, "SAR")
    s20, p_s20 = g(cur, "SMA20"), g(prv, "SMA20")
    s50, s200 = g(cur, "SMA50"), g(cur, "SMA200")
    adx, pdi, ndi = g(cur, "ADX(14)"), g(cur, "+DI"), g(cur, "-DI")
    bias, vr, atr = g(cur, "BIAS(20)"), g(cur, "量比(20)"), g(cur, "ATR(14)")

    if ok(rsi, p_rsi):
        if rsi < 30:                         add(+2, f"RSI={rsi:.1f} 超賣（<30）")
        elif rsi > 70:                       add(-2, f"RSI={rsi:.1f} 超買（>70）")
        elif 30 <= rsi < 50 and rsi > p_rsi: add(+1, f"RSI={rsi:.1f} 低檔回升")
        elif 50 < rsi <= 70 and rsi < p_rsi: add(-1, f"RSI={rsi:.1f} 高檔轉弱")
    if ok(h, p_h):
        if p_h <= 0 < h:          add(+2, f"MACD 柱翻正（{h:.3f}），黃金交叉")
        elif p_h >= 0 > h:        add(-2, f"MACD 柱翻負（{h:.3f}），死亡交叉")
        elif h > 0 and h > p_h:   add(+1, "MACD 柱為正且擴大")
        elif h < 0 and h < p_h:   add(-1, "MACD 柱為負且擴大")
    if ok(k, d, pk, pd_):
        if pk <= pd_ and k > d:   add(+2 if k < 50 else +1, f"KD={k:.1f}/{d:.1f} 黃金交叉" + ("（低檔）" if k < 50 else ""))
        elif pk >= pd_ and k < d: add(-2 if k > 50 else -1, f"KD={k:.1f}/{d:.1f} 死亡交叉" + ("（高檔）" if k > 50 else ""))
        if k < 20 and d < 20:     add(+1, "KD 雙雙低於 20（超賣）")
        elif k > 80 and d > 80:   add(-1, "KD 雙雙高於 80（超買，留意鈍化）")
    if ok(pct_b):
        if pct_b < 0.2:   add(+1, f"%B={pct_b:.2f} 靠近布林下緣")
        elif pct_b > 0.8: add(-1, f"%B={pct_b:.2f} 靠近布林上緣")
    if ok(cci):
        if cci < -100:  add(+1, f"CCI={cci:.0f} 超賣")
        elif cci > 100: add(-1, f"CCI={cci:.0f} 超買")
    if ok(close, sar, p_close, p_sar):
        if p_close <= p_sar and close > sar:   add(+2, f"SAR 翻多（收 {close:.2f} > SAR {sar:.2f}）")
        elif p_close >= p_sar and close < sar: add(-2, f"SAR 翻空（收 {close:.2f} < SAR {sar:.2f}）")
        elif close > sar:                      add(+1, "SAR 多方排列")
        else:                                  add(-1, "SAR 空方排列")
    if ok(s50, s200):
        add(+2, "SMA50 > SMA200 長線多頭") if s50 > s200 else add(-2, "SMA50 < SMA200 長線空頭")
    if ok(close, s20, p_s20):
        if close > s20 and s20 > p_s20:   add(+1, "站上上揚的月線")
        elif close < s20 and s20 < p_s20: add(-1, "跌破下彎的月線")
    if ok(adx, pdi, ndi) and adx > 25:
        add(+2, f"ADX={adx:.1f} 多頭趨勢確立") if pdi > ndi else add(-2, f"ADX={adx:.1f} 空頭趨勢確立")
    if ok(bias):
        if bias > 10:    add(-1, f"BIAS={bias:.1f}% 正乖離過大")
        elif bias < -10: add(+1, f"BIAS={bias:.1f}% 負乖離過大")
    if ok(vr, close, p_close) and vr > 1.5:
        add(+1, f"價漲量增（量比 {vr:.1f}）") if close > p_close else add(-1, f"價跌量增（量比 {vr:.1f}）")

    sb, ss = sum(w for w, _ in buys), -sum(w for w, _ in sells)
    score = sb - ss
    lv = score_to_level(score)
    agree = max(sb, ss) / (sb + ss) if (sb + ss) else 0
    conf = "高" if abs(score) >= 6 and agree >= 0.75 else "中" if abs(score) >= 3 else "低"
    if np.isnan(s200):
        conf += "（長線資料不足）"

    first_close = float(df_disp["收盤價"].iloc[0])
    return {
        "收盤價":   close,
        "漲跌幅":   close / first_close - 1 if first_close else np.nan,
        "分數":     int(score),
        "建議":     lv[1],
        "信心":     conf,
        "停損參考": round(close - 2 * atr, 2) if ok(atr) else None,
        "買進訊號": "\n".join(f"• {t}（+{w}）" for w, t in buys) or "無明顯買進訊號",
        "賣出訊號": "\n".join(f"• {t}（{w}）" for w, t in sells) or "無明顯賣出訊號",
        "關鍵指標": (f"RSI={fmt(rsi)}｜MACD柱={fmt(h, '.3f')}｜KD={fmt(k)}/{fmt(d)}｜"
                    f"%B={fmt(pct_b, '.2f')}｜ADX={fmt(adx)}｜CCI={fmt(cci, '.0f')}｜"
                    f"趨勢={cur.get('Golden_Cross', 'N/A')}"),
    }

# ══════════════════════════════════════════════════════════════════════════════
# 7. Excel 報表（與原版相同）
# ══════════════════════════════════════════════════════════════════════════════
THIN = Side(style="thin", color="BFBFBF")
S = {
    "hdr_font": Font(name="Arial", bold=True, color="FFFFFF", size=11),
    "hdr_fill": PatternFill("solid", start_color="0070C0"),
    "data_font": Font(name="Arial", size=10),
    "small_font": Font(name="Arial", size=9),
    "alt_fill": PatternFill("solid", start_color="EBF3FB"),
    "sum_fill": PatternFill("solid", start_color="D9E1F2"),
    "center": Alignment(horizontal="center", vertical="center"),
    "right": Alignment(horizontal="right", vertical="center"),
    "border": Border(left=THIN, right=THIN, top=THIN, bottom=THIN),
}
FILL_Y = PatternFill("solid", start_color="FFFF00")
FILL_R = PatternFill("solid", start_color="FF4444")
FONT_K = Font(name="Arial", size=9, color="000000")
FONT_W = Font(name="Arial", size=9, color="FFFFFF")

COLOR_RULES = {
    "RSI(14)":  (lambda v: v < 30,   lambda v: v > 70),
    "%B":       (lambda v: v < 0,    lambda v: v > 1),
    "ADX(14)":  (lambda v: v > 25,   lambda v: v < 20),
    "CCI(20)":  (lambda v: v < -100, lambda v: v > 100),
    "BIAS(20)": (lambda v: v < -10,  lambda v: v > 10),
}


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.font, cell.fill = S["hdr_font"], S["hdr_fill"]
        cell.alignment, cell.border = S["center"], S["border"]
    ws.row_dimensions[row].height = 22


def calc_stats(df):
    close = df["收盤價"].astype(float).values
    rets = np.diff(np.log(close)) if len(close) > 1 else np.array([0.0])
    peak = np.maximum.accumulate(close)
    return {
        "年化波動率": float(np.std(rets, ddof=1) * np.sqrt(252)) if len(rets) > 1 else None,
        "最大回撤": float(((close - peak) / peak).min()),
        "hi": float(df["最高價"].max()),
        "lo": float(df["最低價"].min()),
    }


def linear_trend(close):
    x = np.arange(len(close))
    return np.polyval(np.polyfit(x, close, 1), x) if len(close) >= 2 else close


FIB_LEVELS = [("0%", 0), ("23.6%", .236), ("38.2%", .382), ("50%", .5), ("61.8%", .618), ("100%", 1)]


def write_price_sheet(ws, df, ind, title):
    ws.append(["日期", "開盤價", "最高價", "最低價", "收盤價", "成交量", "線性趨勢", "SMA20"])
    style_header(ws)
    close = df["收盤價"].astype(float).values
    trend = linear_trend(close)
    sma20 = ind["SMA20"].values
    for i, row in enumerate(df.itertuples(index=False)):
        r = i + 2
        vals = list(row) + [round(float(trend[i]), 2), to_cell(sma20[i])]
        for ci, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=ci, value=to_cell(v))
            cell.font, cell.border = S["data_font"], S["border"]
            cell.alignment = S["center"] if ci == 1 else S["right"]
            if r % 2 == 0:
                cell.fill = S["alt_fill"]
            cell.number_format = ("yyyy-mm-dd" if ci == 1 else "#,##0" if ci == 6 else "#,##0.00")

    last = ws.max_row
    st_ = calc_stats(df)
    hi, lo = st_["hi"], st_["lo"]
    stats = [
        ("最高收盤價", f"=MAX(E2:E{last})", "#,##0.00"),
        ("最低收盤價", f"=MIN(E2:E{last})", "#,##0.00"),
        ("平均收盤價", f"=AVERAGE(E2:E{last})", "#,##0.00"),
        ("期間漲跌幅", f"=(E{last}-E2)/E2", "0.00%"),
        ("平均成交量", f"=AVERAGE(F2:F{last})", "#,##0"),
        ("年化波動率", st_["年化波動率"], "0.00%"),
        ("最大回撤", st_["最大回撤"], "0.00%"),
    ] + [(f"Fib {lbl}", round(lo + r * (hi - lo), 2), "#,##0.00") for lbl, r in FIB_LEVELS]

    ss = last + 2
    ws.cell(row=ss, column=1, value="統計摘要").font = Font(name="Arial", bold=True, size=11, color="0070C0")
    for i, (lbl, val, nf) in enumerate(stats, start=1):
        lc = ws.cell(row=ss + i, column=1, value=lbl)
        lc.font = Font(name="Arial", bold=True, size=10)
        lc.alignment, lc.border, lc.fill = S["center"], S["border"], S["sum_fill"]
        vc = ws.cell(row=ss + i, column=2, value=val)
        vc.font, vc.alignment, vc.border, vc.number_format = S["data_font"], S["right"], S["border"], nf

    for i, w in enumerate([13, 12, 12, 12, 12, 15, 13, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    ch = LineChart()
    ch.title, ch.style, ch.height, ch.width = f"{title} 收盤價走勢", 10, 15, 28
    ch.y_axis.title, ch.x_axis.title = "價格 (TWD)", "日期"
    ch.x_axis.number_format = "mm-dd"
    for col in (5, 7, 8):
        ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    for sr, color, width, dash in zip(ch.series, ["0070C0", "FF4500", "70AD47"],
                                      [22000, 15000, 15000], [None, "dash", "sysDot"]):
        sr.graphicalProperties.line.solidFill = color
        sr.graphicalProperties.line.width = width
        if dash:
            sr.graphicalProperties.line.dashStyle = dash
        sr.smooth = False
    ws.add_chart(ch, f"D{ss}")


def write_indicator_sheet(ws, ind, title):
    cols = list(ind.columns)
    ws.append(cols)
    style_header(ws)
    for ri, row in enumerate(ind.itertuples(index=False), start=2):
        for ci, (col, raw) in enumerate(zip(cols, row), start=1):
            v = to_cell(raw)
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.border = S["border"]
            cell.alignment = S["center"] if ci == 1 or isinstance(v, str) else S["right"]
            colored = False
            if col in COLOR_RULES and isinstance(v, float):
                yel, red = COLOR_RULES[col]
                if yel(v):
                    cell.fill, cell.font, colored = FILL_Y, FONT_K, True
                elif red(v):
                    cell.fill, cell.font, colored = FILL_R, FONT_W, True
            if not colored:
                cell.font = S["small_font"]
                if ri % 2 == 0:
                    cell.fill = S["alt_fill"]
            if ci == 1:
                cell.number_format = "yyyy-mm-dd"
            elif col == "OBV":
                cell.number_format = "#,##0"
            elif col.startswith("MACD") or col in ("BBW", "%B", "ATR(14)"):
                cell.number_format = "0.0000"
            elif isinstance(v, float):
                cell.number_format = "#,##0.00"
    for ci, col in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = 12 if ci == 1 else min(max(len(col) + 3, 10), 16)
    ws.freeze_panes = "B2"

    last = ws.max_row
    ch = LineChart()
    ch.title, ch.height, ch.width = f"{title} RSI / KD", 10, 28
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 100
    for name in ("RSI(14)", "K(9)", "D(3)"):
        c_idx = cols.index(name) + 1
        ch.add_data(Reference(ws, min_col=c_idx, min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    ch.x_axis.number_format = "mm-dd"
    ws.add_chart(ch, f"B{last + 3}")


def write_advice_sheet(ws, rows, analysis_date):
    headers = ["股票代號", "公司名稱", "收盤價", "區間漲跌幅", "綜合分數", "最新建議",
               "信心程度", "停損參考\n(收盤-2ATR)", "買進訊號說明", "賣出訊號說明", "關鍵指標摘要"]
    widths = [12, 14, 10, 11, 9, 16, 14, 13, 45, 45, 60]
    last_col = get_column_letter(len(headers))

    ws.merge_cells(f"A1:{last_col}1")
    t = ws.cell(row=1, column=1,
                value=f"台股投資建議彙整表｜分析日期：{analysis_date}｜共 {len(rows)} 檔（依綜合分數排序）")
    t.font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    t.fill = PatternFill("solid", start_color="1F497D")
    t.alignment = S["center"]
    ws.row_dimensions[1].height = 32

    ws.merge_cells(f"A2:{last_col}2")
    lg = ws.cell(row=2, column=1, value=(
        "分數 ≥7 強力買進｜≥4 買進｜≥1 觀望偏多｜0 中性｜≤-1 觀望偏空｜≤-4 賣出｜≤-7 強力賣出"
        "　　※ 本表為技術面量化結果，僅供參考，不構成投資建議"))
    lg.font = Font(name="Arial", size=9, italic=True, color="595959")
    lg.alignment = S["center"]
    ws.row_dimensions[2].height = 20

    for ci, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=ci, value=h)
        cell.font, cell.fill, cell.border = S["hdr_font"], S["hdr_fill"], S["border"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 32

    rows = sorted(rows, key=lambda r: r["分數"], reverse=True)
    for ri, r in enumerate(rows, start=4):
        vals = [r["代號"], r["名稱"], r["收盤價"], r["漲跌幅"], r["分數"], r["建議"],
                r["信心"], r["停損參考"], r["買進訊號"], r["賣出訊號"], r["關鍵指標"]]
        for ci, v in enumerate(vals, start=1):
            cell = ws.cell(row=ri, column=ci, value=to_cell(v))
            cell.border = S["border"]
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="center" if ci <= 8 else "left",
                                       vertical="center", wrap_text=True)
            if ri % 2 == 0 and ci != 6:
                cell.fill = PatternFill("solid", start_color="F5F5F5")
        ws.cell(row=ri, column=3).number_format = "#,##0.00"
        ws.cell(row=ri, column=4).number_format = "0.00%"
        ws.cell(row=ri, column=8).number_format = "#,##0.00"

        lv = score_to_level(r["分數"])
        adv = ws.cell(row=ri, column=6)
        adv.fill = PatternFill("solid", start_color=lv[2])
        adv.font = Font(name="Arial", size=10, bold=lv[4], color=lv[3])

        link = ws.cell(row=ri, column=1)
        link.hyperlink = f"#'{r['sheet']}'!A1"
        link.font = Font(name="Arial", size=10, color="0563C1", underline="single")

        n_lines = max(r["買進訊號"].count("\n"), r["賣出訊號"].count("\n")) + 1
        ws.row_dimensions[ri].height = max(30, 15 * n_lines)

    last = 3 + len(rows)
    if rows:
        ws.conditional_formatting.add(f"E4:E{last}", ColorScaleRule(
            start_type="num", start_value=-8, start_color="F8696B",
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=8, end_color="63BE7B"))
        ws.conditional_formatting.add(f"D4:D{last}", ColorScaleRule(
            start_type="min", start_color="F8696B", mid_type="num", mid_value=0,
            mid_color="FFFFFF", end_type="max", end_color="63BE7B"))
    for ci, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:{last_col}{last}"


def build_excel(results, analysis_date):
    wb = Workbook()
    wb.remove(wb.active)
    used = {"投資建議彙整"}
    for r in results:
        sh = safe_title(r["code"], used)
        write_price_sheet(wb.create_sheet(sh), r["df_disp"], r["ind_disp"], r["title"])
        write_indicator_sheet(wb.create_sheet(safe_title(f"{r['code']}_指標", used)),
                              r["ind_disp"], r["title"])
        if r["advice"]:
            r["advice"]["sheet"] = sh
    rows = [r["advice"] for r in results if r["advice"]]
    if rows:
        write_advice_sheet(wb.create_sheet("投資建議彙整", 0), rows, analysis_date)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

# ══════════════════════════════════════════════════════════════════════════════
# 8. 分析流程（st.status 動態進度）
# ══════════════════════════════════════════════════════════════════════════════
def run_analysis(raw_list, display_days, warmup_days):
    t0 = time.time()
    with st.status("🚀 正在分析中…", expanded=True) as status:
        bar = st.progress(0)
        st.write("① 取得台股中文名稱對照表")
        name_map = build_tw_name_map()

        tickers = []
        for x in raw_list:
            t = normalize_ticker(x, name_map)
            if t and t not in tickers:
                tickers.append(t)
        if not tickers:
            status.update(label="股票清單是空的", state="error")
            return None

        bar.progress(10)
        st.write(f"② 查詢 {len(tickers)} 檔公司名稱")
        names = {t: get_company_name(t, name_map) for t in tickers}

        today = datetime.now(TZ).date()
        end = today + timedelta(days=1)
        disp_start = end - timedelta(days=display_days)
        dl_start = end - timedelta(days=display_days + warmup_days)

        bar.progress(20)
        st.write(f"③ 下載股價 {dl_start} ~ {today}")
        data = download_with_fallback(tuple(tickers), str(dl_start), str(end))

        st.write("④ 計算技術指標與評分")
        results, failed, logs = [], [], []
        for i, t in enumerate(tickers):
            bar.progress(30 + int(60 * (i + 1) / len(tickers)))
            code, name = t.split(".")[0], names.get(t, "")
            title = f"{code} {name}".strip()
            if t not in data:
                failed.append(t)
                logs.append(f"❌ {title}：無資料")
                continue
            try:
                full = data[t]
                ind_full = compute_indicators(full)
                mask = full["日期"] >= disp_start
                df_disp = full[mask].reset_index(drop=True)
                ind_disp = ind_full[mask].reset_index(drop=True)
                if df_disp.empty:
                    raise ValueError("顯示區間內沒有資料")
                adv = generate_advice(ind_full, df_disp)
                if adv:
                    adv.update({"代號": t, "名稱": name})
                    logs.append(f"✅ {title}：{len(df_disp)} 筆｜分數 {adv['分數']:+d}｜{adv['建議']}")
                else:
                    logs.append(f"⚠️ {title}：資料太少，無法產生建議")
                results.append({"ticker": t, "code": code, "name": name, "title": title,
                                "df_disp": df_disp, "ind_disp": ind_disp, "advice": adv})
            except Exception as e:
                failed.append(t)
                logs.append(f"❌ {title}：{e}")

        bar.progress(95)
        st.write("⑤ 產生 Excel 報表")
        excel = build_excel(results, today.isoformat()) if results else None
        bar.progress(100)
        elapsed = time.time() - t0
        status.update(label=f"✅ 分析完成：成功 {len(results)} 檔、失敗 {len(failed)} 檔（{elapsed:.1f} 秒）",
                      state="complete", expanded=False)

    st.toast(f"分析完成！共 {len(results)} 檔", icon="🎉")
    return {"tickers": tickers, "results": results, "failed": failed, "logs": logs,
            "excel": excel, "date": today.isoformat(), "elapsed": elapsed}

# ══════════════════════════════════════════════════════════════════════════════
# 9. UI 元件（HTML / Plotly）
# ══════════════════════════════════════════════════════════════════════════════
def md(html_str):
    st.markdown(html_str, unsafe_allow_html=True)


def section(title):
    md(f'<div class="section-title">{title}</div>')


def kpi(label, value, sub="", color=None):
    style = f' style="color:{color}"' if color else ""
    return (f'<div class="card kpi"><div class="lbl">{label}</div>'
            f'<div class="val"{style}>{value}</div><div class="sub">{sub}</div></div>')


def kpi_grid(items):
    md('<div class="kpi-grid">' + "".join(items) + "</div>")


def badge(score):
    lv = score_to_level(score)
    return f'<span class="badge" style="background:#{lv[2]};color:#{lv[3]}">{lv[5]} {lv[1]}</span>'


def svg_spark(vals, color, w=240, h=50):
    v = np.asarray(vals, dtype=float)
    v = v[~np.isnan(v)]
    if len(v) < 2:
        return ""
    lo, hi = v.min(), v.max()
    rng = (hi - lo) or 1
    xs = np.linspace(0, w, len(v))
    ys = h - (v - lo) / rng * (h - 6) - 3
    pts = " ".join(f"{x:.1f},{y:.1f}" for x, y in zip(xs, ys))
    gid = f"g{abs(hash((color, len(v), float(v[-1]), float(v[0])))) % 10**9}"
    return (f'<svg class="spark" viewBox="0 0 {w} {h}" preserveAspectRatio="none">'
            f'<defs><linearGradient id="{gid}" x1="0" y1="0" x2="0" y2="1">'
            f'<stop offset="0%" stop-color="{color}" stop-opacity="0.35"/>'
            f'<stop offset="100%" stop-color="{color}" stop-opacity="0"/></linearGradient></defs>'
            f'<polygon points="0,{h} {pts} {w},{h}" fill="url(#{gid})"/>'
            f'<polyline points="{pts}" fill="none" stroke="{color}" stroke-width="2" '
            f'stroke-linejoin="round" stroke-linecap="round"/></svg>')


def stock_card(r):
    a = r["advice"]
    chg = a["漲跌幅"]
    up = chg >= 0
    color = UP if up else DOWN
    return (f'<div class="card stock-card">'
            f'<div class="sc-top"><div><div class="sc-code">{esc(r["code"])}</div>'
            f'<div class="sc-name">{esc(r["name"])}</div></div>{badge(a["分數"])}</div>'
            f'<div class="sc-price">{a["收盤價"]:,.2f}</div>'
            f'<div class="sc-chg" style="color:{color}">{"▲" if up else "▼"} {chg:+.2%}'
            f'<span style="color:{MUTED};font-weight:400"> 區間</span></div>'
            f'{svg_spark(r["df_disp"]["收盤價"].tolist(), color)}'
            f'<div class="sc-foot"><span>分數 <b>{a["分數"]:+d}</b></span>'
            f'<span>信心 <b>{esc(a["信心"])}</b></span></div></div>')


def signals_html(text, kind):
    title = "📈 買進訊號" if kind == "buy" else "📉 賣出訊號"
    items = []
    for line in text.split("\n"):
        line = line.strip().lstrip("•").strip()
        if not line:
            continue
        m = re.match(r"^(.*)（([+-]?\d+)）$", line)
        if m:
            items.append(f'<li><span>{esc(m.group(1))}</span><span class="w">{m.group(2)}</span></li>')
        else:
            items.append(f'<li class="none"><span>{esc(line)}</span></li>')
    return (f'<div class="card"><div class="sig-title {kind}">{title}</div>'
            f'<ul class="sig {kind}">{"".join(items)}</ul></div>')


def base_layout(fig, height):
    fig.update_layout(
        template="plotly_dark", height=height,
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
        margin=dict(l=10, r=10, t=40, b=10),
        font=dict(family="Noto Sans TC, sans-serif", color="#C9CFE3"),
        legend=dict(orientation="h", yanchor="bottom", y=1.01, x=0, bgcolor="rgba(0,0,0,0)"),
        hoverlabel=dict(bgcolor="#141B2D", bordercolor="#6C8CFF"),
    )
    fig.update_xaxes(showgrid=False, zeroline=False)
    fig.update_yaxes(gridcolor="rgba(255,255,255,0.06)", zeroline=False)
    return fig


def show(fig):
    st.plotly_chart(fig, width="stretch", config={"displayModeBar": False})


def gauge_chart(score):
    lv = score_to_level(score)
    steps = [(-18, -6.5, "rgba(192,0,0,.35)"), (-6.5, -3.5, "rgba(255,112,112,.28)"),
             (-3.5, -0.5, "rgba(255,180,140,.16)"), (-0.5, 0.5, "rgba(242,242,242,.10)"),
             (0.5, 3.5, "rgba(180,230,160,.16)"), (3.5, 6.5, "rgba(146,208,80,.28)"),
             (6.5, 18, "rgba(0,176,80,.35)")]
    fig = go.Figure(go.Indicator(
        mode="gauge+number", value=score,
        number=dict(font=dict(size=46, color="#FFFFFF"), valueformat="+d"),
        title=dict(text=f"綜合分數<br><span style='font-size:13px;color:#{lv[2]}'>{lv[5]} {lv[1]}</span>",
                   font=dict(size=14, color="#AEB6CF")),
        gauge=dict(
            axis=dict(range=[-SCORE_RANGE, SCORE_RANGE], tickwidth=0, tickcolor="#555",
                      tickvals=[-18, -7, -4, 0, 4, 7, 18]),
            bar=dict(color=f"#{lv[2]}", thickness=0.28),
            bgcolor="rgba(0,0,0,0)", borderwidth=0,
            steps=[dict(range=[a, b], color=c) for a, b, c in steps],
        )))
    base_layout(fig, 270)
    fig.update_layout(margin=dict(l=25, r=25, t=60, b=5))
    return fig


def tech_chart(df, ind):
    x = pd.to_datetime(df["日期"])
    all_days = pd.date_range(x.min(), x.max())
    gaps = all_days.difference(x)  # 跳過非交易日

    fig = make_subplots(rows=4, cols=1, shared_xaxes=True, vertical_spacing=0.025,
                        row_heights=[0.52, 0.13, 0.17, 0.18],
                        subplot_titles=("", "成交量", "RSI / KD", "MACD"))
    # 布林通道帶
    fig.add_trace(go.Scatter(x=x, y=ind["BB_Upper"], line=dict(width=0), showlegend=False,
                             hoverinfo="skip"), 1, 1)
    fig.add_trace(go.Scatter(x=x, y=ind["BB_Lower"], line=dict(width=0), fill="tonexty",
                             fillcolor="rgba(108,140,255,0.10)", name="布林通道", hoverinfo="skip"), 1, 1)
    # K 線
    fig.add_trace(go.Candlestick(
        x=x, open=df["開盤價"], high=df["最高價"], low=df["最低價"], close=df["收盤價"], name="K線",
        increasing=dict(line=dict(color=UP), fillcolor=UP),
        decreasing=dict(line=dict(color=DOWN), fillcolor=DOWN)), 1, 1)
    for col, color in [("SMA5", "#FBBF24"), ("SMA20", "#A78BFA"), ("SMA60", "#38BDF8")]:
        fig.add_trace(go.Scatter(x=x, y=ind[col], name=col, line=dict(color=color, width=1.5)), 1, 1)
    fig.add_trace(go.Scatter(x=x, y=ind["SAR"], mode="markers", name="SAR",
                             marker=dict(size=3.5, color="#E5E7EB")), 1, 1)
    # 成交量
    vcol = [UP if c >= o else DOWN for o, c in zip(df["開盤價"], df["收盤價"])]
    fig.add_trace(go.Bar(x=x, y=df["成交量"], marker_color=vcol, opacity=0.75,
                         name="成交量", showlegend=False), 2, 1)
    # RSI / KD
    fig.add_hrect(y0=30, y1=70, fillcolor="rgba(255,255,255,0.03)", line_width=0, row=3, col=1)
    for y, c in [(70, UP), (30, DOWN)]:
        fig.add_hline(y=y, line=dict(color=c, width=1, dash="dot"), row=3, col=1)
    for col, color in [("RSI(14)", "#F472B6"), ("K(9)", "#FBBF24"), ("D(3)", "#38BDF8")]:
        fig.add_trace(go.Scatter(x=x, y=ind[col], name=col, line=dict(color=color, width=1.5)), 3, 1)
    # MACD
    hcol = [UP if v >= 0 else DOWN for v in ind["MACD_Hist"].fillna(0)]
    fig.add_trace(go.Bar(x=x, y=ind["MACD_Hist"], marker_color=hcol, opacity=0.7,
                         name="MACD柱", showlegend=False), 4, 1)
    fig.add_trace(go.Scatter(x=x, y=ind["MACD"], name="DIF", line=dict(color="#FBBF24", width=1.4)), 4, 1)
    fig.add_trace(go.Scatter(x=x, y=ind["MACD_Signal"], name="DEA", line=dict(color="#38BDF8", width=1.4)), 4, 1)

    base_layout(fig, 860)
    fig.update_layout(hovermode="x unified", bargap=0.15)
    fig.update_xaxes(rangeslider_visible=False, rangebreaks=[dict(values=list(gaps))])
    fig.update_yaxes(range=[0, 100], row=3, col=1)
    for ann in fig.layout.annotations:
        ann.update(font=dict(size=12, color=MUTED), x=0, xanchor="left")
    return fig


def score_bar_chart(advs):
    d = sorted(advs, key=lambda a: a["分數"])
    labels = [f'{a["代號"].split(".")[0]} {a["名稱"]}' for a in d]
    colors = [f'#{score_to_level(a["分數"])[2]}' for a in d]
    fig = go.Figure(go.Bar(
        x=[a["分數"] for a in d], y=labels, orientation="h", marker=dict(color=colors),
        text=[f'{a["分數"]:+d}' for a in d], textposition="outside",
        hovertemplate="%{y}<br>分數 %{x:+d}<extra></extra>"))
    base_layout(fig, max(320, 34 * len(d) + 80))
    fig.update_layout(title="綜合分數排行", showlegend=False)
    fig.add_vline(x=0, line=dict(color="rgba(255,255,255,.3)", width=1))
    return fig


def score_scatter(advs):
    fig = go.Figure(go.Scatter(
        x=[a["漲跌幅"] * 100 for a in advs], y=[a["分數"] for a in advs],
        mode="markers+text", text=[a["代號"].split(".")[0] for a in advs],
        textposition="top center", textfont=dict(size=11, color="#C9CFE3"),
        marker=dict(size=16, color=[a["分數"] for a in advs], colorscale="RdYlGn",
                    cmin=-10, cmax=10, line=dict(width=1, color="rgba(255,255,255,.4)")),
        customdata=[a["名稱"] for a in advs],
        hovertemplate="%{text} %{customdata}<br>區間漲跌 %{x:+.2f}%<br>分數 %{y:+d}<extra></extra>"))
    base_layout(fig, 420)
    fig.update_layout(title="動能 × 評分 象限圖", xaxis_title="區間漲跌幅 (%)", yaxis_title="綜合分數")
    fig.add_hline(y=0, line=dict(color="rgba(255,255,255,.25)", dash="dot"))
    fig.add_vline(x=0, line=dict(color="rgba(255,255,255,.25)", dash="dot"))
    return fig


def fib_chart(df):
    x = pd.to_datetime(df["日期"])
    stt = calc_stats(df)
    hi, lo = stt["hi"], stt["lo"]
    fig = go.Figure(go.Scatter(x=x, y=df["收盤價"], name="收盤價", line=dict(color=ACCENT, width=2.2),
                               fill="tozeroy", fillcolor="rgba(108,140,255,.08)"))
    palette = ["#22C55E", "#84CC16", "#FBBF24", "#F59E0B", "#F97316", "#FF4D4F"]
    for (lbl, r), c in zip(FIB_LEVELS, palette):
        y = lo + r * (hi - lo)
        fig.add_hline(y=y, line=dict(color=c, width=1, dash="dot"),
                      annotation_text=f"{lbl}  {y:,.2f}", annotation_position="right",
                      annotation_font=dict(color=c, size=11))
    base_layout(fig, 420)
    pad = (hi - lo) * 0.08
    fig.update_yaxes(range=[lo - pad, hi + pad])
    fig.update_layout(title="費波那契回撤", margin=dict(r=110), showlegend=False)
    return fig


def style_indicator_table(df):
    def colorize(col):
        rule = COLOR_RULES.get(col.name)
        if not rule:
            return [""] * len(col)
        yel, red = rule
        return ["" if pd.isna(v) else
                "background-color:#FACC15;color:#111" if yel(v) else
                "background-color:#EF4444;color:#fff" if red(v) else "" for v in col]
    return df.style.apply(colorize, axis=0).format(precision=2, na_rep="—")

# ══════════════════════════════════════════════════════════════════════════════
# 10. 頁面區塊
# ══════════════════════════════════════════════════════════════════════════════
def render_hero(date=None):
    sub = f"資料日期 {date}｜" if date else ""
    md(f'<div class="hero"><span class="hero-badge">TAIWAN STOCK · TECHNICAL RADAR</span>'
       f'<h1>台股技術分析儀表板</h1>'
       f'<p>{sub}RSI・MACD・KD・布林・ADX・SAR・一目均衡表 加權量化評分，一鍵產生買賣訊號與 Excel 報表</p></div>')


def render_empty():
    feats = [
        ("🎯", "量化評分", "12 類技術指標加權計分，自動分成 7 個多空等級"),
        ("🕯️", "互動 K 線", "K 線、均線、布林、SAR、RSI/KD、MACD 四層聯動圖"),
        ("🃏", "卡片總覽", "每檔股票都有迷你走勢與多空徽章，強弱一目了然"),
        ("📥", "Excel 報表", "一鍵下載含條件格式、圖表與超連結的完整報表"),
    ]
    md('<div class="feat-grid">' + "".join(
        f'<div class="card feat"><div class="ic">{i}</div><h4>{t}</h4><p>{d}</p></div>'
        for i, t, d in feats) + "</div>")
    st.info("👈 在左側完成設定後，按下「開始分析」")


def render_overview(res):
    results = [r for r in res["results"] if r["advice"]]
    advs = [r["advice"] for r in results]

    section("市場總覽")
    avg_score = np.mean([a["分數"] for a in advs]) if advs else 0
    avg_chg = np.mean([a["漲跌幅"] for a in advs]) if advs else 0
    kpi_grid([
        kpi("分析檔數", len(res["tickers"]), f"成功 {len(res['results'])}｜失敗 {len(res['failed'])}"),
        kpi("偏多檔數", sum(a["分數"] >= 1 for a in advs), "分數 ≥ 1", "#4ADE80"),
        kpi("偏空檔數", sum(a["分數"] <= -1 for a in advs), "分數 ≤ -1", "#F87171"),
        kpi("平均分數", f"{avg_score:+.1f}", score_to_level(round(avg_score))[1]),
        kpi("平均區間漲跌", f"{avg_chg:+.2%}", "顯示區間首日 → 最新", UP if avg_chg >= 0 else DOWN),
        kpi("耗時", f"{res['elapsed']:.1f}s", f"分析日 {res['date']}"),
    ])

    c1, c2 = st.columns([3, 1], vertical_alignment="center")
    if res["failed"]:
        c1.warning(f"無法取得資料：{', '.join(res['failed'])}")
    if res["excel"]:
        c2.download_button("📥 下載完整 Excel 報表", data=res["excel"],
                           file_name=f"tw_stocks_{res['date']}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           width="stretch")
    if not advs:
        return

    t1, t2, t3 = st.tabs(["🃏 個股卡片", "📊 分數排行", "📋 彙整表"])
    with t1:
        f1, f2 = st.columns([2, 1])
        flt = f1.segmented_control("篩選", ["全部", "偏多", "中性", "偏空"], default="全部",
                                   label_visibility="collapsed") or "全部"
        sort = f2.selectbox("排序", ["分數高→低", "分數低→高", "漲幅高→低", "漲幅低→高"],
                            label_visibility="collapsed")
        cond = {"全部": lambda s: True, "偏多": lambda s: s >= 1,
                "中性": lambda s: s == 0, "偏空": lambda s: s <= -1}[flt]
        key, rev = {"分數高→低": ("分數", True), "分數低→高": ("分數", False),
                    "漲幅高→低": ("漲跌幅", True), "漲幅低→高": ("漲跌幅", False)}[sort]
        shown = sorted([r for r in results if cond(r["advice"]["分數"])],
                       key=lambda r: r["advice"][key], reverse=rev)
        if shown:
            md('<div class="grid">' + "".join(stock_card(r) for r in shown) + "</div>")
        else:
            st.caption("沒有符合條件的股票")
    with t2:
        a, b = st.columns(2)
        with a:
            show(score_bar_chart(advs))
        with b:
            show(score_scatter(advs))
    with t3:
        rows = sorted(results, key=lambda r: r["advice"]["分數"], reverse=True)
        df = pd.DataFrame([{
            "代號": r["ticker"], "名稱": r["name"],
            "走勢": r["df_disp"]["收盤價"].astype(float).tolist(),
            "收盤價": r["advice"]["收盤價"], "區間漲跌": r["advice"]["漲跌幅"] * 100,
            "分數": r["advice"]["分數"],
            "建議": f'{score_to_level(r["advice"]["分數"])[5]} {r["advice"]["建議"]}',
            "信心": r["advice"]["信心"], "停損參考": r["advice"]["停損參考"],
            "關鍵指標": r["advice"]["關鍵指標"],
        } for r in rows])
        st.dataframe(df, hide_index=True, width="stretch", height=min(600, 38 * len(df) + 40),
                     column_config={
                         "走勢": st.column_config.LineChartColumn("區間走勢", width="medium"),
                         "收盤價": st.column_config.NumberColumn(format="%.2f"),
                         "區間漲跌": st.column_config.NumberColumn(format="%+.2f%%"),
                         "分數": st.column_config.ProgressColumn("綜合分數", min_value=-SCORE_RANGE,
                                                                max_value=SCORE_RANGE, format="%+d"),
                         "停損參考": st.column_config.NumberColumn("停損(收-2ATR)", format="%.2f"),
                         "關鍵指標": st.column_config.TextColumn(width="large"),
                     })
    with st.expander("📝 執行紀錄"):
        st.code("\n".join(res["logs"]), language=None)


def render_detail(res):
    results = res["results"]
    if not results:
        return
    section("個股深度分析")
    opts = {r["title"]: r for r in results}
    r = opts[st.selectbox("選擇個股", list(opts), label_visibility="collapsed")]
    df, ind, a = r["df_disp"], r["ind_disp"], r["advice"]

    if a:
        chg = a["漲跌幅"]
        color = UP if chg >= 0 else DOWN
        md(f'<div class="card detail-head">'
           f'<div><div class="dh-code">{esc(r["ticker"])}</div><div class="dh-name">{esc(r["name"] or r["code"])}</div></div>'
           f'<div><span class="dh-price">{a["收盤價"]:,.2f}</span>'
           f'<span class="dh-chg" style="color:{color}">{"▲" if chg >= 0 else "▼"} {chg:+.2%}</span></div>'
           f'<div class="dh-right">{badge(a["分數"])}<div class="dh-conf">信心程度：{esc(a["信心"])}</div></div></div>')

        last = ind.iloc[-1]
        rsi, k, d = last["RSI(14)"], last["K(9)"], last["D(3)"]
        adx, cci, pb = last["ADX(14)"], last["CCI(20)"], last["%B"]
        state = lambda v, lo, hi, a_, b_, c_: (a_ if v < lo else b_ if v > hi else c_) if pd.notna(v) else "—"

        g, m = st.columns([1, 2])
        with g:
            with st.container(border=True):
                show(gauge_chart(a["分數"]))
        with m:
            kpi_grid([
                kpi("RSI(14)", fmt(rsi), state(rsi, 30, 70, "🟡 超賣", "🔴 超買", "中性區")),
                kpi("KD(9,3)", f"{fmt(k)} / {fmt(d)}",
                    ("K > D 多方" if k > d else "K < D 空方") if pd.notna(k) and pd.notna(d) else "—"),
                kpi("ADX(14)", fmt(adx), state(adx, 20, 25, "盤整", "趨勢成形", "趨勢醞釀")),
                kpi("CCI(20)", fmt(cci, ".0f"), state(cci, -100, 100, "超賣", "超買", "常態")),
                kpi("布林 %B", fmt(pb, ".2f"), state(pb, 0.2, 0.8, "近下緣", "近上緣", "通道中段")),
                kpi("停損參考", fmt(a["停損參考"], ",.2f"), "收盤 − 2×ATR", "#FBBF24"),
            ])
        s1, s2 = st.columns(2)
        with s1:
            md(signals_html(a["買進訊號"], "buy"))
        with s2:
            md(signals_html(a["賣出訊號"], "sell"))
        st.write("")

    tabs = st.tabs(["🕯️ 技術線圖", "📐 統計 & 費波那契", "🧮 指標明細", "📄 原始資料"])
    with tabs[0]:
        show(tech_chart(df, ind))
    with tabs[1]:
        stt = calc_stats(df)
        close = df["收盤價"].astype(float)
        vol = stt["年化波動率"]
        kpi_grid([
            kpi("最高收盤", f"{close.max():,.2f}", color=UP),
            kpi("最低收盤", f"{close.min():,.2f}", color=DOWN),
            kpi("平均收盤", f"{close.mean():,.2f}"),
            kpi("年化波動率", f"{vol:.2%}" if vol is not None else "—"),
            kpi("最大回撤", f"{stt['最大回撤']:.2%}", color=DOWN),
            kpi("平均成交量", f"{df['成交量'].mean():,.0f}"),
        ])
        show(fib_chart(df))
    with tabs[2]:
        st.dataframe(style_indicator_table(ind), hide_index=True, width="stretch", height=480)
        st.caption("🟡 RSI<30｜%B<0｜ADX>25｜CCI<-100｜BIAS<-10　　🔴 RSI>70｜%B>1｜ADX<20｜CCI>100｜BIAS>10")
    with tabs[3]:
        st.dataframe(df, hide_index=True, width="stretch", height=480,
                     column_config={"成交量": st.column_config.NumberColumn(format="localized")})

# ══════════════════════════════════════════════════════════════════════════════
# 11. 主程式
# ══════════════════════════════════════════════════════════════════════════════
def main():
    with st.sidebar:
        md('<div class="brand"><div class="logo">📈</div><div><b>TW Stock Radar</b>'
           '<span>台股技術分析・量化評分</span></div></div>')
        src = st.segmented_control("股票清單來源", ["Google 試算表", "手動輸入"],
                                   default="Google 試算表") or "Google 試算表"
        if src == "Google 試算表":
            sheet_id = st.text_input("SHEET_ID", DEFAULT["SHEET_ID"])
            gid = st.text_input("GID（工作表分頁）", DEFAULT["GID"])
            tcol = st.text_input("代號欄位名稱", DEFAULT["TICKER_COL"])
            st.caption("🔓 試算表需設為「知道連結的任何人皆可檢視」")
        else:
            manual = st.text_area("股票代號", "2330, 2317, 2454, 0050, 6488", height=110,
                                  help="可用逗號、空白或換行分隔；上市/上櫃會自動判斷")

        with st.expander("⚙️ 進階設定"):
            display_days = st.slider("顯示區間（日曆天）", 30, 365, DEFAULT["DISPLAY_DAYS"])
            warmup_days = st.number_input("暖機歷史天數", 250, 1500, DEFAULT["WARMUP_DAYS"], step=50,
                                          help="SMA200 / 一目均衡表需要足夠的歷史資料")
        st.write("")
        run = st.button("🚀 開始分析", type="primary", width="stretch")
        if st.button("🧹 清除快取", width="stretch"):
            st.cache_data.clear()
            st.toast("快取已清除", icon="🧹")
        st.caption("資料來源：Yahoo Finance、臺灣證交所 ISIN")

    res = st.session_state.get("result")
    render_hero(res["date"] if res else None)

    if run:
        if src == "Google 試算表":
            try:
                raw = load_tickers(sheet_id, gid, tcol)
            except Exception as e:
                st.error(f"讀取 Google 試算表失敗：{e}")
                st.stop()
        else:
            raw = parse_manual(manual)
        if not raw:
            st.warning("沒有讀到任何股票代號。")
            st.stop()
        res = run_analysis(raw, int(display_days), int(warmup_days))
        st.session_state["result"] = res

    if not res:
        render_empty()
        return

    render_overview(res)
    render_detail(res)
    md('<div class="footer">※ 本儀表板為技術面量化結果，僅供參考，不構成任何投資建議｜'
       '台股慣例：<span style="color:#FF4D4F">紅漲</span> <span style="color:#22C55E">綠跌</span></div>')


if __name__ == "__main__":
    main()
